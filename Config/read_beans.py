import openpyxl
wb = openpyxl.load_workbook('Config/Datas/__beans__.xlsx')
print('Sheet names:', wb.sheetnames)
ws = wb.active
print('Active sheet:', ws.title)
print('Dimensions:', ws.dimensions)
for row in ws.iter_rows(min_row=1, max_row=20, values_only=True):
    print(row)
